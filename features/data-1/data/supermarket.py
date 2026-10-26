import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook
from PIL import Image, ImageTk
import io
import threading
import urllib.request
import ssl
import os

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
WORKBOOK = os.path.join(_SCRIPT_DIR, "super_market.xlsx")
if not os.path.exists(WORKBOOK):
    WORKBOOK = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", "..", "..", "data", "super_market.xlsx"))

# ---------------------------------------------------------------------------
# Load data from Excel
# ---------------------------------------------------------------------------
book = load_workbook(WORKBOOK, read_only=True, data_only=True)
sheet = book["products"]

headers = [cell.value for cell in sheet[1]]
columns = {str(name).casefold(): idx for idx, name in enumerate(headers)}

products = []
for row in sheet.iter_rows(min_row=2, values_only=True):
    name     = row[columns["englishproductname"]]
    barcode  = row[columns["barcode"]]
    image    = row[columns["image"]]          # URL(s)
    brand    = row[columns.get("brand", 7)]   if columns.get("brand") is not None else ""
    category = row[columns.get("englishcategory", 4)] if columns.get("englishcategory") is not None else ""
    price    = row[columns.get("price", 11)]  if columns.get("price") is not None else ""
    products.append({
        "name":     name or "Unnamed product",
        "barcode":  str(barcode or ""),
        "image":    image or "",
        "brand":    brand or "",
        "category": category or "",
        "price":    price or "",
    })
book.close()

# ---------------------------------------------------------------------------
# Image cache  (URL → PhotoImage)
# ---------------------------------------------------------------------------
image_cache: dict[str, ImageTk.PhotoImage] = {}
IMAGE_SIZE = (220, 220)

# SSL context that skips certificate verification (some CDN certs fail on
# older Windows / Python bundles).
ssl_ctx = ssl.create_default_context()
ssl_ctx.check_hostname = False
ssl_ctx.verify_mode = ssl.CERT_NONE


def _first_url(raw: str) -> str:
    """Return the first http(s) URL found in the cell value."""
    for part in raw.replace("\r\n", "\n").replace("_x000D_", "\n").split("\n"):
        part = part.strip()
        if part.startswith("http"):
            return part
    return ""


def _fetch_image(url: str) -> Image.Image | None:
    """Download an image from *url* and return a PIL Image."""
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=10, context=ssl_ctx) as resp:
            data = resp.read()
        return Image.open(io.BytesIO(data))
    except Exception:
        return None


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------
root = tk.Tk()
root.title("Supermarket Products Viewer")
root.geometry("1100x700")
root.configure(bg="#f0f0f0")

style = ttk.Style(root)
style.theme_use("clam")

# ---- Left panel: search + product list ----
left_frame = ttk.Frame(root, padding=8)
left_frame.pack(side="left", fill="y")

search_var = tk.StringVar()
search_entry = ttk.Entry(left_frame, textvariable=search_var, width=40)
search_entry.pack(fill="x", pady=(0, 6))
search_entry.insert(0, "Search products…")
search_entry.bind("<FocusIn>", lambda e: (
    search_entry.delete(0, tk.END) if search_entry.get() == "Search products…" else None
))

listbox_frame = ttk.Frame(left_frame)
listbox_frame.pack(fill="both", expand=True)

scrollbar = ttk.Scrollbar(listbox_frame, orient="vertical")
product_listbox = tk.Listbox(
    listbox_frame,
    width=48,
    font=("Segoe UI", 10),
    activestyle="none",
    selectbackground="#4a90d9",
    selectforeground="white",
    yscrollcommand=scrollbar.set,
)
scrollbar.config(command=product_listbox.yview)
product_listbox.pack(side="left", fill="both", expand=True)
scrollbar.pack(side="right", fill="y")

# Mapping from listbox index → products index (for search filtering)
visible_indices: list[int] = list(range(len(products)))


def populate_listbox(filter_text: str = ""):
    """Fill the listbox, optionally filtering by *filter_text*."""
    global visible_indices
    product_listbox.delete(0, tk.END)
    visible_indices = []
    ft = filter_text.strip().casefold()
    for i, p in enumerate(products):
        if ft and ft not in p["name"].casefold() and ft not in p["barcode"]:
            continue
        product_listbox.insert(tk.END, f'{p["name"]}  |  {p["barcode"]}')
        visible_indices.append(i)


populate_listbox()


def on_search(*_args):
    txt = search_var.get()
    if txt == "Search products…":
        txt = ""
    populate_listbox(txt)


search_var.trace_add("write", on_search)

# ---- Right panel: product details + image ----
right_frame = ttk.Frame(root, padding=16)
right_frame.pack(side="right", fill="both", expand=True)

# Product info
info_frame = ttk.LabelFrame(right_frame, text="Product Details", padding=12)
info_frame.pack(fill="x")

name_var   = tk.StringVar(value="Select a product")
bar_var    = tk.StringVar()
brand_var  = tk.StringVar()
cat_var    = tk.StringVar()
price_var  = tk.StringVar()

for label_text, var in [
    ("Name:",     name_var),
    ("Barcode:",  bar_var),
    ("Brand:",    brand_var),
    ("Category:", cat_var),
    ("Price:",    price_var),
]:
    row_f = ttk.Frame(info_frame)
    row_f.pack(fill="x", pady=2)
    ttk.Label(row_f, text=label_text, width=10, font=("Segoe UI", 10, "bold")).pack(side="left")
    ttk.Label(row_f, textvariable=var, font=("Segoe UI", 10), wraplength=500).pack(side="left", fill="x")

# Image display
img_frame = ttk.LabelFrame(right_frame, text="Product Image", padding=12)
img_frame.pack(fill="both", expand=True, pady=(12, 0))

img_label = ttk.Label(img_frame, text="No image loaded", anchor="center")
img_label.pack(expand=True)

# Keep a reference so the PhotoImage isn't garbage-collected
_current_photo = None


def show_image_from_url(url: str):
    """Download and display the image (runs the network call in a thread)."""
    global _current_photo

    if not url:
        img_label.config(image="", text="No image available")
        _current_photo = None
        return

    # Check cache first
    if url in image_cache:
        _current_photo = image_cache[url]
        img_label.config(image=_current_photo, text="")
        return

    img_label.config(image="", text="Loading image…")

    def _download():
        global _current_photo
        pil_img = _fetch_image(url)
        if pil_img is None:
            root.after(0, lambda: img_label.config(text="Failed to load image"))
            return
        pil_img.thumbnail(IMAGE_SIZE, Image.LANCZOS)
        photo = ImageTk.PhotoImage(pil_img)
        image_cache[url] = photo

        def _update():
            global _current_photo
            _current_photo = photo
            img_label.config(image=photo, text="")

        root.after(0, _update)

    threading.Thread(target=_download, daemon=True).start()


def on_select(_event=None):
    sel = product_listbox.curselection()
    if not sel:
        return
    idx = visible_indices[sel[0]]
    p = products[idx]
    name_var.set(p["name"])
    bar_var.set(p["barcode"])
    brand_var.set(p["brand"])
    cat_var.set(p["category"])
    price_var.set(str(p["price"]))
    url = _first_url(p["image"])
    show_image_from_url(url)


product_listbox.bind("<<ListboxSelect>>", on_select)

# Select first product on start
if products:
    product_listbox.selection_set(0)
    on_select()

# Status bar
status_lbl = ttk.Label(root, text=f"  {len(products)} products loaded", relief="sunken", anchor="w")
status_lbl.pack(side="bottom", fill="x")

root.mainloop()
